"""Dựng SIXDO_YeuCau06_Dashboard.xlsx: tham số, kế hoạch, nhập liệu tháng × SKU, dashboard có công thức.
Chạy: python build_xlsx.py (trong thư mục yeu-cau-06/do-luong)."""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, '..', 'SIXDO_YeuCau06_Dashboard.xlsx')
plan = json.load(open(os.path.join(HERE, 'plan.json')))

F = 'Arial'
f_norm = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_title = Font(name=F, size=14, bold=True)
f_in = Font(name=F, size=10, color='0000FF')
f_link = Font(name=F, size=10, color='008000')
f_mute = Font(name=F, size=9, italic=True, color='6B6B6B')
fill_in = PatternFill('solid', fgColor='FFF2CC')
fill_head = PatternFill('solid', fgColor='E8E6E1')
fill_ex = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='C9C7C1')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
wrap = Alignment(vertical='top', wrap_text=True)

MONTHS = ['T10', 'T11', 'T12', 'T1', 'T2', 'T3']
MCOLS = ['C', 'D', 'E', 'F', 'G', 'H']
SKUS = ['SP1', 'SP2 xanh', 'SP2 vàng', 'SP3 đen', 'SP3 hồng']
STOCK = [20, 20, 20, 25, 25]

wb = Workbook()

def style_range(ws, rng, font=None, fill=None, border=True, align=None, fmt=None):
    for row in ws[rng]:
        for c in row:
            if font: c.font = font
            if fill: c.fill = fill
            if border: c.border = box
            if align: c.alignment = align
            if fmt: c.number_format = fmt

def base_font(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != F:
                c.font = Font(name=F, size=c.font.size or 10, bold=c.font.bold, italic=c.font.italic, color=c.font.color)

# ======================= HuongDan =======================
ws = wb.active; ws.title = 'HuongDan'
ws.column_dimensions['A'].width = 4; ws.column_dimensions['B'].width = 110
ws['B1'] = 'SIXDO – Dashboard đo lường Amazon US (01/10/2026 – 31/03/2027)'; ws['B1'].font = f_title
lines = [
 ('Mục đích', True),
 ('Theo dõi chuỗi Traffic → CTR → CVR → Orders → Revenue, cùng duy trì doanh số, sell-through 110 units và khách mới. Đi kèm báo cáo Yêu cầu 06.', False),
 ('Cách dùng mỗi tháng (ngày 1–3 của tháng sau)', True),
 ('1. Mở sheet NhapLieu, điền các ô nền vàng chữ xanh cho tháng vừa qua, mỗi SKU một dòng. Nguồn từng cột ghi ở dòng tiêu đề.', False),
 ('2. Mở sheet Dashboard, chọn tháng ở ô C3. Các chỉ số, trạng thái và biểu đồ tự cập nhật.', False),
 ('3. Đối chiếu cột "Trạng thái" với quy tắc quyết định trong báo cáo Yêu cầu 06 mục 6 và Yêu cầu 02 mục 1.4.', False),
 ('Quy ước màu', True),
 ('Chữ xanh trên nền vàng: ô nhập tay. Chữ đen: công thức. Chữ xanh lá: lấy từ sheet khác. Không sửa ô công thức.', False),
 ('Trạng thái: "Đạt" (xanh) · "Theo dõi" (vàng) · "Cảnh báo" (đỏ). Ngưỡng nằm ở sheet ThamSo.', False),
 ('Các sheet', True),
 ('ThamSo: baseline, giá, tồn đầu kỳ, ngưỡng cảnh báo. KeHoach: kế hoạch theo tháng, kịch bản cơ sở (71 units) và mục tiêu (110 units), tính từ Yêu cầu 02.', False),
 ('NhapLieu: số thực tế theo tháng × SKU. Dashboard: chỉ số, so với kế hoạch, trạng thái, biểu đồ, bảng tồn theo SKU.', False),
 ('Lưu ý', True),
 ('Sessions tự nhiên = tổng sessions − click PPC − click ngoài sàn, để không đếm trùng. CVR ngoài sàn đo riêng bằng Amazon Attribution.', False),
 ('Kế hoạch là giả định [GĐ] và sẽ được thay bằng baseline thật sau 07/10/2026 (sheet ThamSo, ô C4).', False),
]
r = 3
for t, h in lines:
    ws[f'B{r}'] = t; ws[f'B{r}'].font = f_bold if h else f_norm; ws[f'B{r}'].alignment = wrap
    r += 1 if not h else 1

# ======================= ThamSo =======================
ts = wb.create_sheet('ThamSo')
ts.column_dimensions['B'].width = 46; ts.column_dimensions['C'].width = 14; ts.column_dimensions['D'].width = 60
ts['B1'] = 'Tham số và ngưỡng'; ts['B1'].font = f_title
ts['B3'], ts['C3'], ts['D3'] = 'Tham số', 'Giá trị', 'Nguồn / ghi chú'
style_range(ts, 'B3:D3', font=f_bold, fill=fill_head, align=center)
params = [
 # key row 4..
 ('Unit Session % baseline', 0.06, '0.0%', '[GĐ] YC02; thay bằng Business Report 90 ngày trước 01/10'),
 ('CVR traffic ngoài sàn', 0.04, '0.0%', '[GĐ] YC02 kịch bản cơ sở'),
 ('Giá niêm yết SP1, SP2 (USD)', 51.99, '#,##0.00', '[XN] đề bài'),
 ('Giá niêm yết SP3 (USD)', 25.99, '#,##0.00', '[XN] đề bài'),
 ('Mục tiêu uplift CVR (lần baseline)', 1.5, '0.00"x"', 'Đề bài: uplift CVR 50%'),
 ('Ngưỡng "Đạt": % kế hoạch units', 0.9, '0%', 'Nội bộ'),
 ('Ngưỡng "Cảnh báo": % kế hoạch units dưới', 0.7, '0%', 'Nội bộ'),
 ('Giá bán TB tối thiểu (% niêm yết)', 0.95, '0%', 'YC02 KPI giữ giá'),
 ('ACoS tối đa SP1, SP2 (hòa vốn B)', 0.461, '0.0%', 'YC02 mục 5'),
 ('ACoS tối đa SP3 (hòa vốn B)', 0.365, '0.0%', 'YC02 mục 5'),
 ('Tỷ lệ hoàn tối đa', 0.20, '0%', 'YC02 KPI chất lượng'),
 ('Rating tối thiểu', 4.3, '0.0', 'YC02 KPI chất lượng'),
 ('CTR tìm kiếm tối thiểu (SQP)', 0.004, '0.00%', '[GĐ] mốc tạm; thay bằng CTR tháng 10 thực tế'),
]
for i, (k, v, fmt, note) in enumerate(params):
    rr = 4 + i
    ts[f'B{rr}'] = k; ts[f'C{rr}'] = v; ts[f'D{rr}'] = note
    ts[f'C{rr}'].number_format = fmt; ts[f'C{rr}'].font = f_in; ts[f'C{rr}'].fill = fill_in
    style_range(ts, f'B{rr}:D{rr}'); ts[f'C{rr}'].font = f_in
P = {k: f'ThamSo!$C${4 + i}' for i, (k, *_ ) in enumerate(params)}
BASE, CVROFF, P12, P3, UPL, OK, WARN, PMIN, AC12, AC3, RET, RAT, CTRMIN = [f'ThamSo!$C${4 + i}' for i in range(13)]

ts['B19'], ts['C19'], ts['D19'] = 'Tồn đầu kỳ theo SKU', 'Units', 'Nguồn'
style_range(ts, 'B19:D19', font=f_bold, fill=fill_head, align=center)
for i, (s, n) in enumerate(zip(SKUS, STOCK)):
    rr = 20 + i
    ts[f'B{rr}'] = s; ts[f'C{rr}'] = n; ts[f'D{rr}'] = '[GĐ] YC02; thay bằng Manage FBA Inventory 01/10'
    style_range(ts, f'B{rr}:D{rr}'); ts[f'C{rr}'].font = f_in; ts[f'C{rr}'].fill = fill_in
ts['B25'] = 'Tổng'; ts['C25'] = '=SUM(C20:C24)'; style_range(ts, 'B25:D25', font=f_bold)
ts['B27'] = 'Danh sách tháng'; ts['B27'].font = f_bold
for i, m in enumerate(MONTHS):
    ts.cell(row=28 + i, column=2, value=m).font = f_norm
MLIST = 'ThamSo!$B$28:$B$33'

# ======================= KeHoach =======================
kh = wb.create_sheet('KeHoach')
kh.column_dimensions['A'].width = 3; kh.column_dimensions['B'].width = 44
for c in MCOLS + ['I']: kh.column_dimensions[c].width = 11
kh.column_dimensions['J'].width = 50
kh['B1'] = 'Kế hoạch theo tháng (tính từ Yêu cầu 02)'; kh['B1'].font = f_title
kh['B2'] = 'Units = (sessions tự nhiên + click PPC) × Unit Session % + click ngoài sàn × CVR ngoài sàn. Ô chữ xanh là giả định [GĐ], sửa được.'
kh['B2'].font = f_mute

def block(start, title, sc):
    d = plan[sc]
    kh[f'B{start}'] = title
    style_range(kh, f'B{start}:J{start}', font=f_bold, fill=fill_head)
    for c, m in zip(MCOLS, MONTHS): kh[f'{c}{start}'] = m; kh[f'{c}{start}'].alignment = center
    kh[f'I{start}'] = 'Tổng'; kh[f'J{start}'] = 'Nguồn'
    rows = [
     ('Sessions tự nhiên', [x['org'] for x in d], True, '#,##0', 'YC02 mục 1.3 (cơ sở 655; mục tiêu 922)'),
     ('Click PPC', [x['ppc'] for x in d], True, '#,##0.0', 'YC02 bảng lịch PPC (234 click)'),
     ('Click ngoài sàn', [x['off'] for x in d], True, '#,##0', 'YC02, YC04 (cơ sở 120; mục tiêu 150)'),
     ('Unit Session % (Amazon)', [x['cvr'] for x in d], True, '0.0%', 'Cơ sở tăng dần 6% → 9%; mục tiêu 9% cả kỳ'),
     ('Chi PPC (USD)', [x['spend'] for x in d], True, '#,##0', 'YC02 mục 9'),
     ('Tỷ trọng units SP1, SP2', [0.70, 0.54, 0.33, 0.73, 0.64, 0.60], True, '0%', 'YC02, cơ cấu SKU theo tháng'),
     ('Tỷ lệ units từ truy vấn theo dịp', [0.20, 0.25, 0.30, 0.40, 0.45, 0.50], True, '0%', 'YC02 KPI: 20% → 50%'),
    ]
    r0 = start + 1
    for j, (lab, vals, inp, fmt, src) in enumerate(rows):
        rr = r0 + j
        kh[f'B{rr}'] = lab; kh[f'J{rr}'] = src
        for c, v in zip(MCOLS, vals):
            kh[f'{c}{rr}'] = v; kh[f'{c}{rr}'].number_format = fmt
            kh[f'{c}{rr}'].font = f_in
        if fmt != '0%' and fmt != '0.0%':
            kh[f'I{rr}'] = f'=SUM(C{rr}:H{rr})'; kh[f'I{rr}'].number_format = fmt
        style_range(kh, f'B{rr}:J{rr}')
        for c in MCOLS: kh[f'{c}{rr}'].font = f_in
    o, pp, off, cvr, sp, sh, occ = [r0 + j for j in range(7)]
    calc = [
     ('Sessions tổng', lambda c: f'={c}{o}+{c}{pp}+{c}{off}', '#,##0'),
     ('Units kế hoạch', lambda c: f'=ROUND(({c}{o}+{c}{pp})*{c}{cvr}+{c}{off}*{CVROFF},0)', '#,##0'),
    ]
    rr = r0 + 7
    for lab, fn, fmt in calc:
        kh[f'B{rr}'] = lab
        for c in MCOLS: kh[f'{c}{rr}'] = fn(c); kh[f'{c}{rr}'].number_format = fmt
        kh[f'I{rr}'] = f'=SUM(C{rr}:H{rr})'; kh[f'I{rr}'].number_format = fmt
        style_range(kh, f'B{rr}:J{rr}'); rr += 1
    u = rr - 1
    kh[f'B{rr}'] = 'Units lũy kế'
    kh[f'C{rr}'] = f'=C{u}'
    for k in range(1, 6): kh[f'{MCOLS[k]}{rr}'] = f'={MCOLS[k-1]}{rr}+{MCOLS[k]}{u}'
    style_range(kh, f'B{rr}:J{rr}', fmt='#,##0'); cum = rr; rr += 1
    kh[f'B{rr}'] = 'Sell-through lũy kế (% của tồn đầu kỳ)'
    for c in MCOLS: kh[f'{c}{rr}'] = f'={c}{cum}/ThamSo!$C$25'
    style_range(kh, f'B{rr}:J{rr}', fmt='0%'); st = rr; rr += 1
    kh[f'B{rr}'] = 'Doanh số theo giá niêm yết (USD)'
    for c in MCOLS: kh[f'{c}{rr}'] = f'={c}{u}*({c}{sh}*{P12}+(1-{c}{sh})*{P3})'
    kh[f'I{rr}'] = f'=SUM(C{rr}:H{rr})'
    style_range(kh, f'B{rr}:J{rr}', fmt='#,##0'); rev = rr; rr += 1
    for row in kh[f'B{r0}:I{rr-1}']:
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('='): c.font = f_norm
    return dict(units=u, cum=cum, st=st, rev=rev, occ=occ, cvr=cvr, spend=sp, end=rr)

KB = block(4, 'Kịch bản cơ sở (71 units)', 'co_so')
KT = block(KB['end'] + 2, 'Kịch bản mục tiêu (110 units)', 'muc_tieu')

# ======================= NhapLieu =======================
nl = wb.create_sheet('NhapLieu')
cols = [
 ('A', 'Tháng', 7, None),
 ('B', 'SKU', 10, None),
 ('C', 'Impression tìm kiếm\n(Search Query Performance)', 15, '#,##0'),
 ('D', 'Click tìm kiếm\n(Search Query Performance)', 13, '#,##0'),
 ('E', 'Sessions tổng\n(Business Report, ASIN con)', 14, '#,##0'),
 ('F', 'Click PPC\n(Advertising console)', 12, '#,##0'),
 ('G', 'Click ngoài sàn\n(Amazon Attribution)', 13, '#,##0'),
 ('H', 'Units bán\n(Business Report)', 11, '#,##0'),
 ('I', 'Doanh số USD\n(Ordered product sales)', 14, '#,##0.00'),
 ('J', 'Chi PPC USD\n(Advertising console)', 12, '#,##0.00'),
 ('K', 'Doanh số PPC USD\n(Advertising console)', 14, '#,##0.00'),
 ('L', 'Units mua qua ngoài sàn\n(Amazon Attribution)', 14, '#,##0'),
 ('M', 'Units hoàn\n(FBA Customer Returns)', 11, '#,##0'),
 ('N', 'Rating cuối tháng\n(trang sản phẩm)', 11, '0.0'),
 ('O', 'Units từ truy vấn theo dịp\n(Search Query Performance)', 15, '#,##0'),
 ('P', 'Units khách mới\n(Repeat Purchase Behavior) [KC]', 15, '#,##0'),
 ('Q', 'Doanh số theo giá niêm yết\n(công thức)', 14, '#,##0.00'),
 ('R', 'Thứ tự tháng\n(công thức)', 9, '0'),
]
for col, name, w, fmt in cols:
    nl.column_dimensions[col].width = w
nl['A1'] = 'Nhập liệu hằng tháng theo SKU'; nl['A1'].font = f_title
nl['A2'] = 'Điền các ô nền vàng chữ xanh. Dòng 4 là VÍ DỤ, không được tính. Cột Q và R là công thức, không sửa.'; nl['A2'].font = f_mute
HR = 3
for col, name, w, fmt in cols:
    nl[f'{col}{HR}'] = name
style_range(nl, f'A{HR}:R{HR}', font=f_bold, fill=fill_head, align=center)
nl.row_dimensions[HR].height = 54
example = ['VÍ DỤ', 'SP2 vàng', 1450, 9, 58, 17, 4, 4, 207.96, 15.00, 103.98, 0, 0, 4.4, 2, 4]
for j, v in enumerate(example):
    c = nl.cell(row=4, column=1 + j, value=v); c.font = f_mute; c.fill = fill_ex; c.border = box
    if cols[j][3]: c.number_format = cols[j][3]
for col in 'QR':
    nl[f'{col}4'].fill = fill_ex; nl[f'{col}4'].border = box
DS, DE = 5, 5 + 30 - 1
r = DS
for m in MONTHS:
    for s in SKUS:
        nl[f'A{r}'] = m; nl[f'B{r}'] = s
        nl[f'Q{r}'] = f'=H{r}*IF(LEFT(B{r},3)="SP3",{P3},{P12})'
        nl[f'R{r}'] = f'=MATCH(A{r},{MLIST},0)'
        for col, name, w, fmt in cols:
            c = nl[f'{col}{r}']; c.border = box
            if fmt: c.number_format = fmt
            if col in 'AB': c.font = f_bold
            elif col in 'QR': c.font = f_norm
            else: c.font = f_in; c.fill = fill_in
        r += 1
nl.freeze_panes = 'C5'
NL = lambda col: f'NhapLieu!${col}${DS}:${col}${DE}'

# ======================= Dashboard =======================
db = wb.create_sheet('Dashboard')
db.column_dimensions['A'].width = 3; db.column_dimensions['B'].width = 40
for c in MCOLS: db.column_dimensions[c].width = 11
db.column_dimensions['I'].width = 11; db.column_dimensions['J'].width = 13; db.column_dimensions['K'].width = 46
db['B1'] = 'SIXDO – Dashboard đo lường'; db['B1'].font = f_title
db['B3'] = 'Tháng đang xem'; db['B3'].font = f_bold
db['C3'] = 'T12'; db['C3'].font = f_in; db['C3'].fill = fill_in; db['C3'].border = box
dv = DataValidation(type='list', formula1=MLIST, allow_blank=False); db.add_data_validation(dv); dv.add('C3')
db['D3'] = f'=MATCH(C3,{MLIST},0)'; db['D3'].font = f_mute
db['E3'] = '← thứ tự tháng (công thức)'; db['E3'].font = f_mute

# --- Thẻ tổng hợp đến tháng đang xem ---
db['B5'] = 'Tổng hợp đến hết tháng đang xem'
style_range(db, 'B5:K5', font=f_bold, fill=fill_head)
db['C5'] = 'Thực tế'; db['D5'] = 'Cơ sở'; db['E5'] = 'Mục tiêu'; db['F5'] = 'Trạng thái'
cumA = f'SUMIFS({NL("H")},{NL("R")},"<="&$D$3)'
cards = [
 ('Units lũy kế', f'={cumA}', f'=INDEX(KeHoach!$C${KB["cum"]}:$H${KB["cum"]},$D$3)', f'=INDEX(KeHoach!$C${KT["cum"]}:$H${KT["cum"]},$D$3)', '#,##0',
  f'=IF(C6=0,"-",IF(C6>=D6*{OK},"Đạt",IF(C6>=D6*{WARN},"Theo dõi","Cảnh báo")))', 'So với kịch bản cơ sở'),
 ('Sell-through lũy kế', '=C6/ThamSo!$C$25', f'=INDEX(KeHoach!$C${KB["st"]}:$H${KB["st"]},$D$3)', f'=INDEX(KeHoach!$C${KT["st"]}:$H${KT["st"]},$D$3)', '0%',
  '=F6', 'Đi cùng units lũy kế'),
 ('Unit Session % lũy kế (sessions Amazon)', f'=IF(SUMIFS({NL("E")},{NL("R")},"<="&$D$3)-SUMIFS({NL("G")},{NL("R")},"<="&$D$3)=0,0,(C6-SUMIFS({NL("L")},{NL("R")},"<="&$D$3))/(SUMIFS({NL("E")},{NL("R")},"<="&$D$3)-SUMIFS({NL("G")},{NL("R")},"<="&$D$3)))',
  f'=INDEX(KeHoach!$C${KB["cvr"]}:$H${KB["cvr"]},$D$3)', f'=INDEX(KeHoach!$C${KT["cvr"]}:$H${KT["cvr"]},$D$3)', '0.0%',
  f'=IF(C8=0,"-",IF(C8>={BASE}*{UPL},"Đạt",IF(C8>={BASE},"Theo dõi","Cảnh báo")))', 'Đạt khi ≥ 1,5 lần baseline'),
 ('Uplift CVR (lần baseline)', f'=C8/{BASE}', f'=D8/{BASE}', f'=E8/{BASE}', '0.00"x"', '=F8', 'Đề bài: ≥ 1,5x'),
 ('Doanh số lũy kế (USD)', f'=SUMIFS({NL("I")},{NL("R")},"<="&$D$3)', f'=SUMPRODUCT((COLUMN(KeHoach!$C${KB["rev"]}:$H${KB["rev"]})-2<=$D$3)*KeHoach!$C${KB["rev"]}:$H${KB["rev"]})',
  f'=SUMPRODUCT((COLUMN(KeHoach!$C${KT["rev"]}:$H${KT["rev"]})-2<=$D$3)*KeHoach!$C${KT["rev"]}:$H${KT["rev"]})', '#,##0',
  f'=IF(C10=0,"-",IF(C10>=D10*{OK},"Đạt",IF(C10>=D10*{WARN},"Theo dõi","Cảnh báo")))', 'Kế hoạch tính theo giá niêm yết'),
 ('Giá bán TB (% giá niêm yết)', f'=IF(SUMIFS({NL("Q")},{NL("R")},"<="&$D$3)=0,0,C10/SUMIFS({NL("Q")},{NL("R")},"<="&$D$3))', f'={PMIN}', f'={PMIN}', '0%',
  f'=IF(C11=0,"-",IF(C11>={PMIN},"Đạt","Cảnh báo"))', 'Giữ giá: ≥ 95%'),
]
for i, (lab, a, b, c_, fmt, stt, note) in enumerate(cards):
    rr = 6 + i
    db[f'B{rr}'] = lab; db[f'C{rr}'] = a; db[f'D{rr}'] = b; db[f'E{rr}'] = c_; db[f'F{rr}'] = stt; db[f'K{rr}'] = note
    for col in 'CDE': db[f'{col}{rr}'].number_format = fmt
    style_range(db, f'B{rr}:K{rr}')
    db[f'F{rr}'].alignment = center
    for col in 'DE': db[f'{col}{rr}'].font = f_link

# --- Bảng theo tháng ---
T0 = 14
db[f'B{T0}'] = 'Theo tháng'
for c, m in zip(MCOLS, MONTHS): db[f'{c}{T0}'] = m
db[f'I{T0}'] = 'Cả kỳ'; db[f'J{T0}'] = 'Trạng thái tháng đang xem'; db[f'K{T0}'] = 'Công thức / nguồn'
style_range(db, f'B{T0}:K{T0}', font=f_bold, fill=fill_head, align=center)
S = lambda col, c: f'SUMIFS({NL(col)},{NL("A")},{c}${T0})'
mrows = [
 ('Impression tìm kiếm', lambda c: f'={S("C", c)}', '#,##0', 'sum', None, 'SQP'),
 ('CTR tìm kiếm', lambda c: f'=IF({c}15=0,0,{S("D", c)}/{c}15)', '0.00%', 'ratio:{D}/15', 'ctr', 'Click ÷ impression (SQP)'),
 ('Sessions tổng', lambda c: f'={S("E", c)}', '#,##0', 'sum', None, 'Business Report'),
 ('  Sessions tự nhiên', lambda c: f'={c}17-{c}19-{c}20', '#,##0', 'sum', None, 'Tổng − PPC − ngoài sàn'),
 ('  Click PPC', lambda c: f'={S("F", c)}', '#,##0', 'sum', None, 'Advertising console'),
 ('  Click ngoài sàn', lambda c: f'={S("G", c)}', '#,##0', 'sum', None, 'Amazon Attribution'),
 ('Unit Session % (sessions Amazon)', lambda c: f'=IF({c}17-{c}20=0,0,({c}23-{S("L", c)})/({c}17-{c}20))', '0.0%', 'cvr', 'cvr', '(Units − units ngoài sàn) ÷ (sessions − click ngoài sàn)'),
 ('Uplift CVR (lần baseline)', lambda c: f'={c}21/{BASE}', '0.00"x"', 'upl', 'upl', 'Mục tiêu ≥ 1,5x'),
 ('Units bán', lambda c: f'={S("H", c)}', '#,##0', 'sum', None, 'Business Report'),
 ('Units kế hoạch (cơ sở)', lambda c: f'=KeHoach!{c}{KB["units"]}', '#,##0', 'sum', None, 'Sheet KeHoach'),
 ('% kế hoạch', lambda c: f'=IF({c}24=0,0,{c}23/{c}24)', '0%', 'ratio:23/24', 'plan', 'Đạt ≥ 90%; cảnh báo < 70%'),
 ('Sell-through lũy kế', lambda c: f'=SUMIFS({NL("H")},{NL("R")},"<="&MATCH({c}${T0},{MLIST},0))/ThamSo!$C$25', '0%', 'last', None, 'Units lũy kế ÷ tồn đầu kỳ'),
 ('Doanh số (USD)', lambda c: f'={S("I", c)}', '#,##0', 'sum', None, 'Ordered product sales'),
 ('Giá bán TB (% niêm yết)', lambda c: f'=IF({S("Q", c)}=0,0,{c}27/{S("Q", c)})', '0%', 'price', 'price', 'Doanh số ÷ units × giá niêm yết'),
 ('Chi PPC (USD)', lambda c: f'={S("J", c)}', '#,##0', 'sum', None, 'Advertising console'),
 ('ACoS', lambda c: f'=IF({S("K", c)}=0,0,{c}29/{S("K", c)})', '0.0%', 'acos', 'acos', 'Chi PPC ÷ doanh số PPC; so với hòa vốn B'),
 ('TACoS', lambda c: f'=IF({c}27=0,0,{c}29/{c}27)', '0.0%', 'ratio:29/27', None, 'Chi PPC ÷ tổng doanh số'),
 ('Tỷ lệ hoàn', lambda c: f'=IF({c}23=0,0,{S("M", c)}/{c}23)', '0.0%', 'ret', 'ret', 'Units hoàn ÷ units bán; tối đa 20%'),
 ('Rating thấp nhất trong các SKU', lambda c: f'=IF({c}23=0,0,_xlfn.MINIFS({NL("N")},{NL("A")},{c}${T0},{NL("N")},">0"))', '0.0', 'min', 'rat', 'Tối thiểu 4,3'),
 ('% units từ truy vấn theo dịp', lambda c: f'=IF({c}23=0,0,{S("O", c)}/{c}23)', '0%', 'occ', 'occ', 'Mục tiêu 20% (T10) → 50% (T3)'),
 ('% units từ khách mới [KC]', lambda c: f'=IF({c}23=0,0,{S("P", c)}/{c}23)', '0%', 'ratio:P', None, 'Repeat Purchase Behavior'),
 ('Units qua ngoài sàn', lambda c: f'={S("L", c)}', '#,##0', 'sum', None, 'Amazon Attribution'),
]
for i, (lab, fn, fmt, agg, stt, note) in enumerate(mrows):
    rr = T0 + 1 + i
    db[f'B{rr}'] = lab; db[f'K{rr}'] = note
    for c in MCOLS: db[f'{c}{rr}'] = fn(c); db[f'{c}{rr}'].number_format = fmt
    if agg == 'sum': db[f'I{rr}'] = f'=SUM(C{rr}:H{rr})'
    elif agg == 'last': db[f'I{rr}'] = f'=H{rr}'
    elif agg == 'cvr': db[f'I{rr}'] = f'=IF(I17-I20=0,0,(I23-I36)/(I17-I20))'
    elif agg == 'upl': db[f'I{rr}'] = f'=I21/{BASE}'
    elif agg == 'price': db[f'I{rr}'] = f'=IF(SUM({NL("Q")})=0,0,I27/SUM({NL("Q")}))'
    elif agg == 'acos': db[f'I{rr}'] = f'=IF(SUM({NL("K")})=0,0,I29/SUM({NL("K")}))'
    elif agg == 'ret': db[f'I{rr}'] = f'=IF(I23=0,0,SUM({NL("M")})/I23)'
    elif agg == 'min': db[f'I{rr}'] = f'=_xlfn.MINIFS(C{rr}:H{rr},C{rr}:H{rr},">0")'
    elif agg == 'occ': db[f'I{rr}'] = f'=IF(I23=0,0,SUM({NL("O")})/I23)'
    elif agg == 'ratio:P': db[f'I{rr}'] = f'=IF(I23=0,0,SUM({NL("P")})/I23)'
    elif agg == 'ratio:23/24': db[f'I{rr}'] = '=IF(I24=0,0,I23/I24)'
    elif agg == 'ratio:29/27': db[f'I{rr}'] = '=IF(I27=0,0,I29/I27)'
    elif agg == 'ratio:{D}/15': db[f'I{rr}'] = f'=IF(I15=0,0,SUM({NL("D")})/I15)'
    db[f'I{rr}'].number_format = fmt
    sel = f'INDEX(C{rr}:H{rr},$D$3)'
    if stt:
        rules = {
         'ctr': f'=IF({sel}=0,"-",IF({sel}>={CTRMIN},"Đạt","Theo dõi"))',
         'cvr': f'=IF({sel}=0,"-",IF({sel}>={BASE}*{UPL},"Đạt",IF({sel}>={BASE},"Theo dõi","Cảnh báo")))',
         'upl': f'=IF({sel}=0,"-",IF({sel}>={UPL},"Đạt",IF({sel}>=1,"Theo dõi","Cảnh báo")))',
         'plan': f'=IF({sel}=0,"-",IF({sel}>={OK},"Đạt",IF({sel}>={WARN},"Theo dõi","Cảnh báo")))',
         'price': f'=IF({sel}=0,"-",IF({sel}>={PMIN},"Đạt","Cảnh báo"))',
         'acos': f'=IF({sel}=0,"-",IF({sel}<={AC3},"Đạt",IF({sel}<={AC12},"Theo dõi","Cảnh báo")))',
         'ret': f'=IF(INDEX(C23:H23,$D$3)=0,"-",IF({sel}<={RET}*0.75,"Đạt",IF({sel}<={RET},"Theo dõi","Cảnh báo")))',
         'rat': f'=IF({sel}=0,"-",IF({sel}>={RAT},"Đạt","Cảnh báo"))',
         'occ': f'=IF(INDEX(C23:H23,$D$3)=0,"-",IF({sel}>=INDEX(KeHoach!$C${KB["occ"]}:$H${KB["occ"]},$D$3),"Đạt","Theo dõi"))',
        }
        db[f'J{rr}'] = rules[stt]; db[f'J{rr}'].alignment = center
    style_range(db, f'B{rr}:K{rr}')
    if lab.startswith('  '): db[f'B{rr}'].font = Font(name=F, size=10, color='52514E')
    db[f'C{rr}'].number_format = fmt
LAST = T0 + len(mrows)

# --- Bảng tồn theo SKU ---
K0 = LAST + 3
db[f'B{K0}'] = 'Tồn kho theo SKU đến hết tháng đang xem'
hdr = ['SKU', 'Tồn đầu kỳ', 'Units bán lũy kế', 'Tồn còn lại', 'Sell-through', 'Units tháng đang xem', 'Tháng còn hàng (ước tính)', 'Việc cần làm']
for j, h in enumerate(hdr): db.cell(row=K0 + 1, column=2 + j, value=h)
style_range(db, f'B{K0+1}:I{K0+1}', font=f_bold, fill=fill_head, align=center)
db.merge_cells(f'I{K0+1}:K{K0+1}')
for i, s in enumerate(SKUS):
    rr = K0 + 2 + i
    db[f'B{rr}'] = s
    db[f'C{rr}'] = f'=ThamSo!C{20 + i}'; db[f'C{rr}'].font = f_link
    db[f'D{rr}'] = f'=SUMIFS({NL("H")},{NL("B")},B{rr},{NL("R")},"<="&$D$3)'
    db[f'E{rr}'] = f'=C{rr}-D{rr}'
    db[f'F{rr}'] = f'=IF(C{rr}=0,0,D{rr}/C{rr})'; db[f'F{rr}'].number_format = '0%'
    db[f'G{rr}'] = f'=SUMIFS({NL("H")},{NL("B")},B{rr},{NL("A")},$C$3)'
    db[f'H{rr}'] = f'=IF(G{rr}=0,"-",E{rr}/G{rr})'; db[f'H{rr}'].number_format = '0.0'
    db[f'I{rr}'] = f'=IF(E{rr}<=1,"Dừng quảng cáo biến thể còn ≤ 1 unit",IF(AND($D$3>=5,E{rr}>2*G{rr}*(7-$D$3)),"Tồn cao: bundle, coupon theo dịp tháng 3","Theo kế hoạch"))'
    db.merge_cells(f'I{rr}:K{rr}')
    style_range(db, f'B{rr}:K{rr}')
    db[f'C{rr}'].font = f_link

# Định dạng trạng thái
green = PatternFill('solid', fgColor='D9F2D9'); yellow = PatternFill('solid', fgColor='FFF0C2'); red = PatternFill('solid', fgColor='F8D4D4')
for rng in (f'F6:F11', f'J{T0+1}:J{LAST}'):
    db.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Đạt"'], fill=green))
    db.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Theo dõi"'], fill=yellow))
    db.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Cảnh báo"'], fill=red))
db.conditional_formatting.add(f'I{K0+2}:I{K0+6}', FormulaRule(formula=[f'LEFT(I{K0+2},4)="Dừng"'], fill=red))
db.conditional_formatting.add(f'I{K0+2}:I{K0+6}', FormulaRule(formula=[f'LEFT(I{K0+2},3)="Tồn"'], fill=yellow))

# Hàng phụ cho biểu đồ: units thực tế, kế hoạch cơ sở, mục tiêu; sell-through
C0 = K0 + 9
db[f'B{C0}'] = 'Dữ liệu cho biểu đồ'; db[f'B{C0}'].font = f_bold
for c, m in zip(MCOLS, MONTHS): db[f'{c}{C0}'] = m
db[f'B{C0+1}'] = 'Units thực tế'; db[f'B{C0+2}'] = 'Units kế hoạch cơ sở'; db[f'B{C0+3}'] = 'Units kế hoạch mục tiêu'
db[f'B{C0+4}'] = 'Sell-through thực tế'; db[f'B{C0+5}'] = 'Sell-through cơ sở'; db[f'B{C0+6}'] = 'Sell-through mục tiêu'
for c in MCOLS:
    db[f'{c}{C0+1}'] = f'={c}23'; db[f'{c}{C0+2}'] = f'=KeHoach!{c}{KB["units"]}'; db[f'{c}{C0+3}'] = f'=KeHoach!{c}{KT["units"]}'
    db[f'{c}{C0+4}'] = f'={c}26'; db[f'{c}{C0+5}'] = f'=KeHoach!{c}{KB["st"]}'; db[f'{c}{C0+6}'] = f'=KeHoach!{c}{KT["st"]}'
    for k in (4, 5, 6): db[f'{c}{C0+k}'].number_format = '0%'
style_range(db, f'B{C0}:H{C0+6}')

bc = BarChart(); bc.type = 'col'; bc.title = 'Units theo tháng: thực tế và kế hoạch'; bc.y_axis.title = 'Units'
bc.add_data(Reference(db, min_col=2, max_col=8, min_row=C0 + 1, max_row=C0 + 3), from_rows=True, titles_from_data=True)
bc.set_categories(Reference(db, min_col=3, max_col=8, min_row=C0))
bc.height, bc.width = 7.5, 16
for s, col in zip(bc.series, ('2A78D6', 'B7D3F6', 'EB6834')):
    s.graphicalProperties.solidFill = col; s.graphicalProperties.line.solidFill = col
db.add_chart(bc, f'M5')
lc = LineChart(); lc.title = 'Sell-through lũy kế'; lc.y_axis.number_format = '0%'
lc.add_data(Reference(db, min_col=2, max_col=8, min_row=C0 + 4, max_row=C0 + 6), from_rows=True, titles_from_data=True)
lc.set_categories(Reference(db, min_col=3, max_col=8, min_row=C0))
lc.height, lc.width = 7.5, 16
for s, col in zip(lc.series, ('2A78D6', '86B6EF', 'EB6834')):
    s.graphicalProperties.line.solidFill = col; s.graphicalProperties.line.width = 22000
db.add_chart(lc, f'M22')
db.freeze_panes = 'C5'

for sh in wb.worksheets:
    for row in sh.iter_rows():
        for c in row:
            if c.font.name != F:
                c.font = Font(name=F, size=c.font.size, bold=c.font.bold, italic=c.font.italic, color=c.font.color)
wb.calculation.fullCalcOnLoad = True
wb.active = wb.index(db)
wb.save(OUT)
print('saved', OUT, 'rows', T0, LAST, K0, C0)
